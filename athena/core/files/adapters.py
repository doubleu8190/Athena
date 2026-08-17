"""内置文件适配器 — 文档、数据、图片、源码和压缩包的提取与分析。

本模块提供 7 个内置适配器，覆盖 Athena 系统支持的所有文件类型：

- ``TextAdapter``  — 纯文本、Markdown、JSON、YAML、CSV 等
- ``PdfAdapter``   — PDF 文档（含 OCR 回退和表格提取）
- ``WordAdapter``  — Word (.docx) 文档
- ``ExcelAdapter`` — Excel (.xlsx/.xlsm) 电子表格
- ``ImageAdapter`` — 图片文件（OCR 文本提取）
- ``CodeAdapter``  — 源代码文件（符号索引和依赖分析）
- ``ArchiveAdapter`` — 压缩包（递归提取内部文件）

所有适配器实现 ``FileAdapter`` 协议，返回 ``ExtractionResult``。
"""

from __future__ import annotations

import ast
import asyncio
import csv
import importlib
import io
import re
import stat
import tarfile
import zipfile
from collections import Counter
from pathlib import Path, PurePosixPath
from typing import Any

from athena.config.settings import Settings
from athena.core.files.base import ExtractedUnit, ExtractionResult
from athena.models.file import AdapterInfo
from athena.utils.logging import get_logger

logger = get_logger(__name__)

# 支持的文本文件扩展名集合
TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".json", ".yaml", ".yml", ".toml", ".xml", ".html",
    ".css", ".sql", ".log", ".ini", ".cfg",
}

# 支持的源代码扩展名 → 语言名称映射
CODE_EXTENSIONS = {
    ".py": "python", ".js": "javascript", ".jsx": "javascript",
    ".ts": "typescript", ".tsx": "typescript", ".java": "java",
    ".go": "go", ".rs": "rust", ".c": "c", ".h": "c",
    ".cc": "cpp", ".cpp": "cpp", ".hpp": "cpp", ".cs": "c_sharp",
}


def _decode(data: bytes) -> str:
    """将字节数据解码为字符串。

    优先使用 charset_normalizer 自动检测编码，失败时回退到 UTF-8
    （errors='replace' 确保不会抛出异常）。

    Args:
        data: 原始字节数据。

    Returns:
        解码后的字符串。
    """
    try:
        from charset_normalizer import from_bytes

        match = from_bytes(data).best()
        if match is not None:
            return str(match)
    except Exception as exc:
        logger.warning("charset_detection_failed", error=str(exc))
    return data.decode("utf-8", errors="replace")


def _rapidocr_text(img_content: str | bytes | Path) -> str:
    """使用 RapidOCR 从支持的输入类型中提取文本。"""
    from rapidocr_onnxruntime import RapidOCR

    result, _ = RapidOCR()(img_content)
    return "\n".join(item[1] for item in (result or []))


def _pil_image_png_bytes(image: Any) -> bytes:
    """将 PIL 图像编码为 RapidOCR 支持的图片字节。"""
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


class TextAdapter:
    """纯文本和结构化文本文件适配器。

    支持的格式：纯文本 (.txt)、Markdown (.md)、JSON、YAML、XML、HTML、
    CSS、SQL、日志文件、INI/CFG 配置文件、CSV 等。

    CSV 文件会额外提取表格数据（最多 200 行），并支持 pandas 分析。

    Capabilities: read, search, summarize, analyze
    """

    info = AdapterInfo(
        name="text", version="2.0", mime_types=["text/plain", "text/markdown", "application/json"],
        extensions=sorted(TEXT_EXTENSIONS | {".csv"}), capabilities=["read", "search", "summarize", "analyze"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取文本文件内容（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取文本内容和 CSV 表格数据。"""
        text = _decode(path.read_bytes())
        metadata: dict[str, Any] = {"line_count": len(text.splitlines())}
        tables: list[dict[str, Any]] = []
        # CSV 文件额外解析表格结构
        if path.suffix.lower() == ".csv":
            rows = list(csv.reader(io.StringIO(text)))
            if rows:
                metadata.update({"row_count": max(0, len(rows) - 1), "columns": rows[0]})
                tables.append({"name": path.name, "headers": rows[0], "rows": rows[1:201]})
        return ExtractionResult(
            units=[ExtractedUnit(text, {"path": path.name, "start_line": 1}, {"kind": "text"})],
            metadata=metadata, tables=tables,
        )

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析文本文件：非 CSV 返回基本统计，CSV 使用 pandas 进行列级分析。"""
        text = _decode(path.read_bytes())
        if path.suffix.lower() != ".csv":
            return {"characters": len(text), "lines": len(text.splitlines())}
        try:
            import pandas as pd

            frame = await asyncio.to_thread(pd.read_csv, path)
            return _frame_analysis(frame)
        except Exception as exc:
            return {"error": str(exc)}


class PdfAdapter:
    """PDF 文档适配器。

    提取策略：先用 pypdf 提取文本层，对空白页自动回退到 OCR
    （pypdfium2 + RapidOCR）。表格提取使用 pdfplumber。

    Capabilities: read, search, summarize, analyze, extract_table
    """

    info = AdapterInfo(
        name="pdf", version="2.0", mime_types=["application/pdf"], extensions=[".pdf"],
        capabilities=["read", "search", "summarize", "analyze", "extract_table"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取 PDF 内容（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取 PDF 文本、OCR 和表格数据。"""
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        units: list[ExtractedUnit] = []
        empty_pages: list[int] = []
        # 逐页提取文本，空白页自动 OCR 回退
        for index, page in enumerate(reader.pages, 1):
            content = (page.extract_text() or "").strip()
            if not content:
                empty_pages.append(index)
                content = self._ocr_page(path, index - 1)
            if content:
                units.append(ExtractedUnit(content, {"page": index}, {"kind": "page", "ocr": index in empty_pages}))
        # 使用 pdfplumber 提取结构化表格
        tables: list[dict[str, Any]] = []
        try:
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                for page_no, page in enumerate(pdf.pages, 1):
                    for table_no, table in enumerate(page.extract_tables(), 1):
                        tables.append({"page": page_no, "table": table_no, "rows": table})
        except Exception as exc:
            logger.warning("pdf_table_extract_failed", path=str(path), error=str(exc))
        return ExtractionResult(
            units=units,
            metadata={"pages": len(reader.pages), "ocr_pages": empty_pages, "table_count": len(tables)},
            tables=tables,
        )

    @staticmethod
    def _ocr_page(path: Path, page_index: int) -> str:
        """对单页 PDF 执行 OCR 文本识别。

        使用 pypdfium2 渲染页面为图片（2x 缩放），再通过 RapidOCR 提取文本。
        任何异常均记录日志并返回空字符串（降级为无文本页）。

        Args:
            path: PDF 文件路径。
            page_index: 页码索引（从 0 开始）。

        Returns:
            OCR 识别的文本内容，失败时返回空字符串。
        """
        try:
            import pypdfium2 as pdfium

            document = pdfium.PdfDocument(str(path))
            image = document[page_index].render(scale=2).to_pil()
            return _rapidocr_text(_pil_image_png_bytes(image))
        except Exception as exc:
            logger.warning("pdf_ocr_failed", path=str(path), page_index=page_index, error=str(exc))
            return ""

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析 PDF：返回页数、OCR 页列表和表格数量等元数据。"""
        result = await asyncio.to_thread(self._extract, path)
        return result.metadata


class WordAdapter:
    """Word 文档适配器 (.docx)。

    按段落提取文本内容（保留段落样式信息），并提取嵌入表格。
    空段落自动过滤。

    Capabilities: read, search, summarize, analyze, extract_table
    """

    info = AdapterInfo(
        name="word", version="2.0", mime_types=["application/vnd.openxmlformats-officedocument.wordprocessingml.document"],
        extensions=[".docx"], capabilities=["read", "search", "summarize", "analyze", "extract_table"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取 Word 文档内容（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取段落文本和嵌入表格。"""
        from docx import Document

        doc = Document(str(path))
        # 按段落提取文本，附带段落号和样式信息
        units = [
            ExtractedUnit(p.text, {"paragraph": i + 1}, {"style": p.style.name if p.style else ""})
            for i, p in enumerate(doc.paragraphs) if p.text.strip()
        ]
        # 提取嵌入表格
        tables = []
        for index, table in enumerate(doc.tables, 1):
            tables.append({"table": index, "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        return ExtractionResult(units=units, metadata={"paragraphs": len(doc.paragraphs), "tables": len(tables)}, tables=tables)

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析 Word 文档：返回段落数和表格数量等元数据。"""
        result = await asyncio.to_thread(self._extract, path)
        return result.metadata


class ExcelAdapter:
    """Excel 电子表格适配器 (.xlsx/.xlsm)。

    按工作表提取数据，保留公式信息。每个工作表生成一个内容单元
    和一个表格记录。分析模式使用 pandas 进行列级统计。

    Capabilities: read, search, summarize, analyze, extract_table
    """

    info = AdapterInfo(
        name="excel", version="2.0",
        mime_types=["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel.sheet.macroEnabled.12"],
        extensions=[".xlsx", ".xlsm"], capabilities=["read", "search", "summarize", "analyze", "extract_table"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取 Excel 内容（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取所有工作表的数据和公式。"""
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        units: list[ExtractedUnit] = []
        tables: list[dict[str, Any]] = []
        sheet_meta: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            lines: list[str] = []
            formulas: list[dict[str, Any]] = []
            for row in sheet.iter_rows(values_only=False):
                values = [cell.value for cell in row]
                # 收集公式单元格
                formulas.extend(
                    {"cell": cell.coordinate, "formula": cell.value}
                    for cell in row if cell.data_type == "f"
                )
                if any(value is not None for value in values):
                    rows.append(values)
                    lines.append("\t".join("" if value is None else str(value) for value in values))
            units.append(ExtractedUnit(
                "\n".join(lines),
                {"sheet": sheet.title, "start_row": 1, "end_row": sheet.max_row},
                {"kind": "sheet", "formulas": formulas},
            ))
            tables.append({"sheet": sheet.title, "rows": rows[:201]})
            sheet_meta.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column, "formulas": formulas})
        workbook.close()
        return ExtractionResult(units=units, metadata={"sheets": sheet_meta}, tables=tables)

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析 Excel：使用 pandas 对每个工作表进行列级统计分析。"""
        try:
            import pandas as pd

            sheets = await asyncio.to_thread(pd.read_excel, path, sheet_name=None)
            return {name: _frame_analysis(frame) for name, frame in sheets.items()}
        except Exception as exc:
            return {"error": str(exc)}


class ImageAdapter:
    """图片文件适配器。

    使用 RapidOCR 提取图片中的文本内容（OCR）。提取结果包含
    图片尺寸、色彩模式等元数据。视觉分析（描述图片内容）需配合
    支持 vision 的 LLM 模型，在 runtime 层实现。

    Capabilities: read, search, summarize, analyze
    """

    info = AdapterInfo(
        name="image", version="2.0", mime_types=["image/png", "image/jpeg", "image/webp", "image/tiff"],
        extensions=[".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"],
        capabilities=["read", "search", "summarize", "analyze"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取图片 OCR 文本（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取图片元数据和 OCR 文本。"""
        from PIL import Image

        with Image.open(path) as image:
            width, height, mode = image.width, image.height, image.mode
        text = ""
        try:
            text = _rapidocr_text(path)
        except Exception as exc:
            logger.warning("image_ocr_failed", path=str(path), error=str(exc))
        return ExtractionResult(
            units=[ExtractedUnit(text, {"image": path.name}, {"kind": "ocr"})] if text else [],
            metadata={"width": width, "height": height, "mode": mode, "ocr_available": bool(text)},
        )

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析图片：返回尺寸、模式和 OCR 可用性等元数据。"""
        result = await asyncio.to_thread(self._extract, path)
        return {**result.metadata, "note": "视觉模型未配置时仅提供 OCR 结果"}


class CodeAdapter:
    """源代码文件适配器。

    支持 Python、JavaScript、TypeScript、Java、Go、Rust、C/C++、C# 等语言。
    提取源码内容的同时构建符号索引（类、函数、接口等）和依赖关系图。
    Python 使用 AST 解析，其他语言优先使用 tree-sitter，回退到正则匹配。

    Capabilities: read, search, summarize, analyze, symbols, references, call_graph
    """

    info = AdapterInfo(
        name="code", version="2.0", mime_types=["text/x-python", "application/javascript", "text/x-java-source"],
        extensions=sorted(CODE_EXTENSIONS),
        capabilities=["read", "search", "summarize", "analyze", "symbols", "references", "call_graph"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取代码文件内容和符号索引（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract_one, path, path.name)

    def _extract_one(self, path: Path, relative_path: str) -> ExtractionResult:
        """同步提取源码、符号表和依赖关系。"""
        text = _decode(path.read_bytes())
        language = CODE_EXTENSIONS.get(path.suffix.lower(), "text")
        symbols, dependencies = _code_index(text, language, relative_path)
        return ExtractionResult(
            units=[ExtractedUnit(text, {"path": relative_path, "start_line": 1}, {"language": language, "kind": "source"})],
            metadata={"language": language, "lines": len(text.splitlines())}, symbols=symbols, dependencies=dependencies,
        )

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析代码文件：返回语言分布和符号数量。"""
        result = await asyncio.to_thread(self._extract_one, path, path.name)
        return {"languages": {result.metadata.get("language", "text"): 1}, "symbols": len(result.symbols)}


class ArchiveAdapter:
    """压缩包适配器 (.zip/.tar/.tgz/.tar.gz)。

    安全解压后递归提取内部文件，对代码文件构建符号索引和依赖图，
    对文本文件提取内容。内置多层安全校验：路径穿越检测、压缩比炸弹
    防护、符号链接拒绝、嵌套压缩包拒绝。

    Capabilities: read, search, summarize, analyze, symbols, references, call_graph
    """

    info = AdapterInfo(
        name="archive", version="2.0", mime_types=["application/zip", "application/x-tar", "application/gzip"],
        extensions=[".zip", ".tar", ".tgz", ".tar.gz"],
        capabilities=["read", "search", "summarize", "analyze", "symbols", "references", "call_graph"],
    )

    async def extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """提取压缩包内容（在线程池中执行 I/O）。"""
        return await asyncio.to_thread(self._extract, path, settings, workspace)

    def _extract(self, path: Path, settings: Settings, workspace: Path) -> ExtractionResult:
        """安全解压并递归提取内部文件的代码符号和文本内容。"""
        extracted = workspace / "archive"
        extracted.mkdir(parents=True, exist_ok=True)
        members = self._safe_extract(path, extracted, settings)
        code = CodeAdapter()
        text_adapter = TextAdapter()
        units: list[ExtractedUnit] = []
        symbols: list[dict[str, Any]] = []
        dependencies: list[dict[str, Any]] = []
        languages: Counter[str] = Counter()
        # 按文件类型分发到对应适配器
        for relative in members:
            item = extracted / relative
            suffix = item.suffix.lower()
            if suffix in CODE_EXTENSIONS:
                result = code._extract_one(item, relative)
                units.extend(result.units)
                symbols.extend(result.symbols)
                dependencies.extend(result.dependencies)
                languages[result.metadata["language"]] += 1
            elif suffix in TEXT_EXTENSIONS or suffix == ".csv":
                result = text_adapter._extract(item)
                for unit in result.units:
                    unit.locator["path"] = relative
                units.extend(result.units)
        return ExtractionResult(
            units=units, symbols=symbols, dependencies=dependencies,
            metadata={"files": len(members), "indexed_files": len(units), "languages": dict(languages)},
        )

    def _safe_extract(self, path: Path, destination: Path, settings: Settings) -> list[str]:
        """安全解压压缩包到目标目录。

        校验流程：收集成员记录 → 安全性校验 → 逐个解压。
        支持 ZIP 和 TAR 格式，其他格式抛出 ValueError。

        Args:
            path: 压缩包文件路径。
            destination: 解压目标目录。
            settings: 全局配置（含安全限制参数）。

        Returns:
            成功解压的文件相对路径列表。

        Raises:
            ValueError: 不支持的格式、路径穿越、压缩比炸弹等安全违规。
        """
        records: list[tuple[str, int, int, bool, Any]] = []
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as archive:
                for member in archive.infolist():
                    is_link = stat.S_ISLNK(member.external_attr >> 16)
                    records.append((member.filename, member.file_size, member.compress_size, is_link, member))
                self._validate_archive(records, settings)
                for name, _, _, _, member in records:
                    target = self._safe_target(destination, name)
                    if member.is_dir():
                        target.mkdir(parents=True, exist_ok=True)
                    else:
                        target.parent.mkdir(parents=True, exist_ok=True)
                        with archive.open(member) as source, target.open("wb") as output:
                            output.write(source.read())
        elif tarfile.is_tarfile(path):
            with tarfile.open(path) as archive:
                for member in archive.getmembers():
                    records.append((member.name, member.size, member.size, member.issym() or member.islnk(), member))
                self._validate_archive(records, settings, compressed_size=path.stat().st_size)
                for name, _, _, _, member in records:
                    target = self._safe_target(destination, name)
                    if member.isdir():
                        target.mkdir(parents=True, exist_ok=True)
                    elif member.isfile():
                        target.parent.mkdir(parents=True, exist_ok=True)
                        source = archive.extractfile(member)
                        if source:
                            with source, target.open("wb") as output:
                                output.write(source.read())
        else:
            raise ValueError("不支持的压缩包格式")
        return [name for name, _, _, _, _ in records if (destination / name).is_file()]

    @staticmethod
    def _safe_target(root: Path, name: str) -> Path:
        """计算安全的解压目标路径，防止路径穿越攻击。

        Args:
            root: 解压根目录。
            name: 压缩包内的条目路径。

        Returns:
            解析后的安全绝对路径。

        Raises:
            ValueError: 路径包含绝对路径组件或 ``..`` 穿越。
        """
        pure = PurePosixPath(name)
        if pure.is_absolute() or ".." in pure.parts:
            raise ValueError("压缩包包含路径穿越条目")
        target = (root / Path(*pure.parts)).resolve()
        if root.resolve() not in target.parents:
            raise ValueError("压缩包条目逃逸工作目录")
        return target

    @staticmethod
    def _validate_archive(
        records: list[tuple[str, int, int, bool, Any]],
        settings: Settings,
        compressed_size: int | None = None,
    ) -> None:
        """校验压缩包安全性（条目数、大小、压缩比、符号链接、嵌套）。

        Args:
            records: 条目记录列表，每项为 (名称, 原始大小, 压缩大小, 是否链接, 原始对象)。
            settings: 全局配置，包含各项安全限制阈值。
            compressed_size: 压缩包整体大小（TAR 格式使用），ZIP 格式为 ``None``。

        Raises:
            ValueError: 任一安全校验失败。
        """
        if len(records) > settings.file_archive_max_entries:
            raise ValueError("压缩包文件数量超过安全上限")
        total = sum(size for _, size, _, _, _ in records)
        compressed = compressed_size or sum(max(1, size) for _, _, size, _, _ in records)
        if total > settings.file_archive_max_uncompressed_bytes:
            raise ValueError("压缩包解压大小超过安全上限")
        if total / max(1, compressed) > settings.file_archive_max_ratio:
            raise ValueError("压缩包压缩比超过安全上限")
        if any(is_link for _, _, _, is_link, _ in records):
            raise ValueError("压缩包不允许符号链接")
        archive_suffixes = {".zip", ".tar", ".tgz", ".gz"}
        if any(Path(name).suffix.lower() in archive_suffixes for name, _, _, _, _ in records):
            raise ValueError("不递归解压嵌套压缩包")

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析压缩包：返回提示信息（实际分析使用已生成的符号和依赖索引）。"""
        return {"task": task, "note": "代码项目分析使用已生成的符号和依赖索引"}


def _frame_analysis(frame: Any) -> dict[str, Any]:
    """对 pandas DataFrame 执行列级统计分析。

    返回行数、列名、数据类型、缺失值统计和数值列的描述性统计。

    Args:
        frame: pandas DataFrame 对象。

    Returns:
        包含 rows/columns/dtypes/missing/description 字段的分析结果字典。
    """
    numeric = frame.select_dtypes(include="number")
    description = numeric.describe().replace({float("nan"): None}).to_dict() if not numeric.empty else {}
    return {
        "rows": int(len(frame)), "columns": [str(column) for column in frame.columns],
        "dtypes": {str(k): str(v) for k, v in frame.dtypes.items()},
        "missing": {str(k): int(v) for k, v in frame.isna().sum().items()},
        "description": description,
    }


def _code_index(text: str, language: str, path: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """构建代码文件的符号索引和依赖关系图。

    解析策略优先级：
    1. Python → AST 解析（最准确）
    2. 其他语言 → tree-sitter 解析（语法精确）
    3. 回退 → 正则匹配（兼容性最好，精度最低）

    Args:
        text: 源代码文本。
        language: 编程语言标识（如 ``"python"``、``"typescript"``）。
        path: 文件相对路径，用于符号定位。

    Returns:
        (symbols, dependencies) 二元组：符号列表和依赖关系列表。
    """
    if language == "python":
        try:
            return _python_index(text, path)
        except SyntaxError as exc:
            logger.warning("python_ast_index_failed", path=path, error=str(exc), fallback="regex")
    if language != "python":
        parsed = _tree_sitter_index(text, language, path)
        if parsed is not None:
            return parsed
    # 正则回退：匹配类/函数/结构体等定义
    symbols: list[dict[str, Any]] = []
    dependencies: list[dict[str, Any]] = []
    pattern = re.compile(r"^\s*(?:class|interface|struct|enum|def|function|func|fn)\s+([A-Za-z_$][\w$]*)", re.MULTILINE)
    lines = text.splitlines()
    for match in pattern.finditer(text):
        line = text.count("\n", 0, match.start()) + 1
        symbols.append({"path": path, "language": language, "name": match.group(1), "qualified_name": match.group(1),
                        "kind": "symbol", "start_line": line, "end_line": line, "signature": lines[line - 1][:500]})
    # 匹配 import/require/include 语句
    for match in re.finditer(r"(?:import|from|require\s*\(|#include\s*[<\"])([^\n;\)\">]+)", text):
        dependencies.append({"source": path, "target": match.group(1).strip(), "kind": "import", "metadata": {}})
    return symbols, dependencies


def _tree_sitter_index(text: str, language: str, path: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]] | None:
    """Use tree-sitter when the optional grammar pack is available.

    Grammar downloads or parser changes must never make uploads fail, so all
    failures intentionally return ``None`` and the caller uses text fallback.
    """
    try:
        from tree_sitter_language_pack import get_parser

        parser = get_parser(language)
        tree = parser.parse(text.encode("utf-8"))
        symbols: list[dict[str, Any]] = []
        dependencies: list[dict[str, Any]] = []
        interesting = {
            "function_definition", "function_declaration", "method_definition",
            "class_definition", "class_declaration", "struct_declaration",
            "interface_declaration", "enum_declaration", "function_item",
        }
        stack = [tree.root_node]
        while stack:
            node = stack.pop()
            if node.type in interesting:
                name_node = node.child_by_field_name("name")
                if name_node is not None:
                    name = text[name_node.start_byte:name_node.end_byte]
                    symbols.append({
                        "path": path, "language": language, "name": name,
                        "qualified_name": name, "kind": node.type,
                        "start_line": node.start_point[0] + 1,
                        "end_line": node.end_point[0] + 1,
                        "signature": text.splitlines()[node.start_point[0]][:500],
                    })
            stack.extend(reversed(node.children))
        for match in re.finditer(r"(?:import|from|require\s*\(|#include\s*[<\"])([^\n;\)\">]+)", text):
            dependencies.append({"source": path, "target": match.group(1).strip(), "kind": "import", "metadata": {}})
        return symbols, dependencies
    except Exception as exc:
        logger.warning("tree_sitter_index_failed", path=path, language=language, error=str(exc), fallback="regex")
        return None


def _python_index(text: str, path: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """使用 Python AST 构建符号索引和依赖关系图。

    遍历 AST 节点提取类定义、函数定义（含异步函数）、import 语句
    和函数调用关系。类/函数支持嵌套，通过 ``parents`` 栈维护限定名。

    Args:
        text: Python 源代码文本。
        path: 文件相对路径，用于符号定位。

    Returns:
        (symbols, dependencies) 二元组。

    Raises:
        SyntaxError: 源码语法错误时由 ``ast.parse()`` 抛出，
            调用方应回退到正则匹配。
    """
    tree = ast.parse(text)
    symbols: list[dict[str, Any]] = []
    dependencies: list[dict[str, Any]] = []
    parents: list[str] = []  # 嵌套类/函数的名称栈，用于构建限定名

    class Visitor(ast.NodeVisitor):
        def visit_ClassDef(self, node: ast.ClassDef) -> None:
            qualified = ".".join([*parents, node.name])
            symbols.append({"path": path, "language": "python", "name": node.name, "qualified_name": qualified,
                            "kind": "class", "start_line": node.lineno, "end_line": getattr(node, "end_lineno", node.lineno),
                            "signature": f"class {node.name}"})
            parents.append(node.name)
            self.generic_visit(node)
            parents.pop()

        def visit_FunctionDef(self, node: ast.FunctionDef) -> None:
            self._visit_function(node)

        def visit_AsyncFunctionDef(self, node: ast.AsyncFunctionDef) -> None:
            self._visit_function(node)

        def _visit_function(self, node: ast.FunctionDef | ast.AsyncFunctionDef) -> None:
            qualified = ".".join([*parents, node.name])
            args = ", ".join(arg.arg for arg in node.args.args)
            symbols.append({"path": path, "language": "python", "name": node.name, "qualified_name": qualified,
                            "kind": "function", "start_line": node.lineno, "end_line": getattr(node, "end_lineno", node.lineno),
                            "signature": f"def {node.name}({args})"})
            parents.append(node.name)
            self.generic_visit(node)
            parents.pop()

        def visit_Import(self, node: ast.Import) -> None:
            for name in node.names:
                dependencies.append({"source": path, "target": name.name, "kind": "import", "metadata": {"line": node.lineno}})

        def visit_ImportFrom(self, node: ast.ImportFrom) -> None:
            dependencies.append({"source": path, "target": node.module or "", "kind": "import", "metadata": {"line": node.lineno}})

        def visit_Call(self, node: ast.Call) -> None:
            target = ""
            if isinstance(node.func, ast.Name):
                target = node.func.id
            elif isinstance(node.func, ast.Attribute):
                target = node.func.attr
            if target:
                dependencies.append({"source": ".".join(parents) or path, "target": target, "kind": "call", "metadata": {"line": node.lineno}})
            self.generic_visit(node)

    Visitor().visit(tree)
    return symbols, dependencies
