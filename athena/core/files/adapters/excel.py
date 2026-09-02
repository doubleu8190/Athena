"""Excel 文件适配器。"""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from athena.config.settings import Settings
from athena.core.files.base import ExtractedUnit, ExtractionContext, ExtractionResult
from athena.models.file import AdapterInfo

from .common import _frame_analysis


class ExcelAdapter:
    """Excel 电子表格适配器，支持工作表、公式和表格数据提取。"""

    info = AdapterInfo(
        name="excel",
        version="2.0",
        mime_types=[
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        ],
        extensions=[".xlsx", ".xlsm"],
        capabilities=["read", "search", "summarize", "analyze", "extract_table"],
    )

    async def extract(
        self, context: ExtractionContext, settings: Settings
    ) -> ExtractionResult:
        """提取 Excel 内容。"""
        return await asyncio.to_thread(self._extract, context.path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取所有工作表的数据和公式。"""
        from openpyxl import load_workbook

        workbook = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
        units: list[ExtractedUnit] = []
        tables: list[dict[str, Any]] = []
        sheet_meta: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            lines: list[str] = []
            formulas: list[dict[str, Any]] = []
            for row in sheet.iter_rows(values_only=False):
                values = [cell.value for cell in row]
                formulas.extend(
                    {"cell": cell.coordinate, "formula": cell.value}
                    for cell in row
                    if cell.data_type == "f"
                )
                if any(value is not None for value in values):
                    rows.append(values)
                    lines.append(
                        "\t".join(
                            "" if value is None else str(value) for value in values
                        )
                    )
            units.append(
                ExtractedUnit(
                    "\n".join(lines),
                    {"sheet": sheet.title, "start_row": 1, "end_row": sheet.max_row},
                    {"kind": "sheet", "formulas": formulas},
                )
            )
            tables.append({"sheet": sheet.title, "rows": rows[:201]})
            sheet_meta.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "formulas": formulas,
                }
            )
        workbook.close()
        return ExtractionResult(
            units=units, metadata={"sheets": sheet_meta}, tables=tables
        )

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析 Excel 并返回每个工作表的列级统计。"""
        try:
            import pandas as pd

            sheets = await asyncio.to_thread(pd.read_excel, path, sheet_name=None)
            return {name: _frame_analysis(frame) for name, frame in sheets.items()}
        except Exception as exc:
            return {"error": str(exc)}
