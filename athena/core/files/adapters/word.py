"""Word 文件适配器。"""

from __future__ import annotations

import asyncio
import io
from pathlib import Path
from typing import Any
from docx import Document
from athena.config.settings import Settings
from athena.core.files.base import ExtractedUnit, ExtractionContext, ExtractionResult
from athena.models.file import AdapterInfo

from .common import _pil_image_png_bytes, _rapidocr_text, logger


def _word_image_units(document: Any) -> list[ExtractedUnit]:
    """提取 DOCX 文档及其关联部件中的嵌入图片和 OCR 文本。"""
    from PIL import Image

    image_parts: list[Any] = []
    seen_names: set[str] = set()
    pending_parts = [document.part]
    visited_parts: set[int] = set()
    while pending_parts:
        part = pending_parts.pop()
        if id(part) in visited_parts:
            continue
        visited_parts.add(id(part))
        if getattr(part, "content_type", "").startswith("image/"):
            part_name = str(part.partname)
            if part_name not in seen_names:
                seen_names.add(part_name)
                image_parts.append(part)
        pending_parts.extend(getattr(part, "related_parts", {}).values())

    units: list[ExtractedUnit] = []
    for index, part in enumerate(image_parts, 1):
        image_name = Path(str(part.partname)).name
        width: int | None = None
        height: int | None = None
        ocr_text = ""
        try:
            with Image.open(io.BytesIO(part.blob)) as image:
                width, height = image.size
                ocr_text = _rapidocr_text(_pil_image_png_bytes(image))
        except Exception as exc:
            logger.warning(
                "word_image_extract_failed", image=image_name, error=str(exc)
            )
        units.append(
            ExtractedUnit(
                ocr_text or f"[Embedded image: {image_name}]",
                {"image": image_name, "image_index": index},
                {
                    "kind": "image",
                    "format": part.content_type.removeprefix("image/"),
                    "width": width,
                    "height": height,
                    "ocr_available": bool(ocr_text),
                },
            )
        )
    return units


class WordAdapter:
    """Word 文档适配器，支持段落、表格和嵌入图片。"""

    info = AdapterInfo(
        name="word",
        version="2.1",
        mime_types=[
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ],
        extensions=[".docx"],
        capabilities=["read", "search", "summarize", "analyze", "extract_table"],
    )

    async def extract(
        self, context: ExtractionContext, settings: Settings
    ) -> ExtractionResult:
        """提取 Word 文档内容。"""
        return await asyncio.to_thread(self._extract, context.path)

    def _extract(self, path: Path) -> ExtractionResult:
        """同步提取段落、嵌入表格和图片。"""

        doc = Document(str(path))
        units = [
            ExtractedUnit(
                p.text, {"paragraph": i}, {"style": p.style.name if p.style else ""}
            )
            for i, p in enumerate(doc.paragraphs, 1)
            if p.text.strip()
        ]
        tables = [
            {
                "table": index,
                "rows": [[cell.text for cell in row.cells] for row in table.rows],
            }
            for index, table in enumerate(doc.tables, 1)
        ]
        image_units = _word_image_units(doc)
        units.extend(image_units)
        return ExtractionResult(
            units=units,
            metadata={
                "paragraphs": len(doc.paragraphs),
                "tables": len(tables),
                "images": len(image_units),
            },
            tables=tables,
        )

    async def analyze(self, path: Path, task: str) -> dict[str, Any]:
        """分析 Word 文档并返回文件元数据。"""
        result = await asyncio.to_thread(self._extract, path)
        return result.metadata
