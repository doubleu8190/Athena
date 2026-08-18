from __future__ import annotations

from datetime import datetime
import sys
import types
import zipfile

import pytest

from athena.config.settings import Settings
from athena.core.files.adapters import ExcelAdapter, ImageAdapter, PdfAdapter, WordAdapter, ArchiveAdapter
from athena.core.files.base import ExtractionContext
from athena.core.files.capabilities import register_file_capabilities
from athena.core.files.runtime import FileAccessError, FileIntelligenceRuntime
from athena.core.files.storage import FileTooLargeError, StorageLayer
from athena.core.files.tasks import FileTaskWorker
from athena.core.tools.manager import UnifiedToolManager
from athena.db.database import Database
from athena.main import _persist_registered_tool_configs
from athena.models import Message, MessageRole
from athena.models.file import FileChunk, FileTaskStatus, FileTaskType


class _FakeLLM:
    async def ainvoke(self, _messages):
        class Response:
            content = "summary"

        return Response()


async def _chunks(*values: bytes):
    for value in values:
        yield value


def _context(path, workspace, filename: str | None = None, mime_type: str = "") -> ExtractionContext:
    return ExtractionContext(path=path, workspace=workspace, filename=filename or path.name, mime_type=mime_type)


@pytest.mark.asyncio
async def test_content_addressed_storage_deduplicates_and_limits(tmp_path):
    storage = StorageLayer(tmp_path, max_upload_bytes=8)
    first = await storage.save_stream(_chunks(b"same"))
    second = await storage.save_stream(_chunks(b"sa", b"me"))
    assert first.storage_key == second.storage_key
    assert len(storage.blob_keys()) == 1

    with pytest.raises(FileTooLargeError):
        await storage.save_stream(_chunks(b"123456789"))


@pytest.mark.asyncio
async def test_parse_search_message_binding_and_session_isolation(tmp_path):
    db = Database(str(tmp_path / "files.db"))
    await db.connect()
    try:
        await db.sessions.create("one")
        await db.sessions.create("two")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "files.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        blob = await runtime.storage.save_stream(_chunks(b"alpha beta\nsecond line\n"))
        attachment = await db.files.create_attachment(
            session_id="one", filename="notes.txt", mime_type="text/plain",
            size_bytes=blob.size_bytes, sha256=blob.sha256, storage_key=blob.storage_key,
        )
        parsed = await runtime.parse_attachment(attachment.id)
        await runtime.index_attachment(attachment.id)
        result = await runtime.search_file("one", attachment.id, "alpha")
        assert parsed["chunk_count"] == 1
        assert result["results"][0]["locator"]["path"] == "notes.txt"

        with pytest.raises(FileAccessError):
            await runtime.read_file("two", attachment.id)
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_parse_csv_preserves_suffix_for_table_artifact(tmp_path):
    db = Database(str(tmp_path / "csv-artifacts.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "csv-artifacts.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        blob = await runtime.storage.save_stream(_chunks(b"name,value\nalpha,1\nbeta,2\n"))
        attachment = await db.files.create_attachment(
            session_id="session", filename="data.csv", mime_type="text/csv",
            size_bytes=blob.size_bytes, sha256=blob.sha256, storage_key=blob.storage_key,
        )

        parsed = await runtime.parse_attachment(attachment.id)
        key = FileIntelligenceRuntime.cache_key(
            attachment, "tables", {}, "2.0", settings.primary_llm.model
        )
        artifact = await db.files.get_artifact(key)

        assert parsed["row_count"] == 2
        assert parsed["columns"] == ["name", "value"]
        assert artifact is not None
        assert artifact["metadata"]["count"] == 1

        tables = await runtime.extract_table("session", attachment.id)
        assert tables["tables"][0]["headers"] == ["name", "value"]
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_queue_recovery_and_active_task_idempotency(tmp_path):
    db = Database(str(tmp_path / "queue.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        attachment = await db.files.create_attachment(
            session_id="session", filename="a.txt", mime_type="text/plain",
            size_bytes=1, sha256="0" * 64, storage_key="blobs/00/placeholder",
        )
        first = await db.files.create_task("session", attachment.id, FileTaskType.FILE_PARSE)
        duplicate = await db.files.create_task("session", attachment.id, FileTaskType.FILE_PARSE)
        assert duplicate.id == first.id
        claimed = await db.files.claim_next_task()
        assert claimed and claimed.status == FileTaskStatus.RUNNING
        assert await db.files.recover_running_tasks() == 1
        recovered = await db.files.get_task(first.id)
        assert recovered and recovered.status == FileTaskStatus.QUEUED
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_soft_delete_removes_parsed_file_data(tmp_path):
    db = Database(str(tmp_path / "delete.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        message = Message(
            id="message",
            session_id="session",
            role=MessageRole.USER,
            content="use file",
            timestamp=datetime.now(),
        )
        await db.messages.save(message)
        attachment = await db.files.create_attachment(
            session_id="session", filename="a.txt", mime_type="text/plain",
            size_bytes=1, sha256="2" * 64, storage_key="blobs/22/placeholder",
        )
        await db.files.bind_message("session", message.id, [attachment.id])
        await db.files.replace_chunks(
            attachment.id,
            [FileChunk(id="chunk", attachment_id=attachment.id, ordinal=0, content="secret text")],
        )
        await db.files.put_artifact(attachment.id, "summary", "cache", "secret summary")
        await db.files.replace_code_index(
            attachment.id,
            [{
                "path": "a.py", "language": "python", "name": "f",
                "qualified_name": "f", "kind": "function",
                "start_line": 1, "end_line": 1, "signature": "def f()",
            }],
            [],
        )

        assert await db.files.soft_delete_attachment(attachment.id, "session") is True

        assert await db.files.get_chunks(attachment.id) == []
        assert await db.files.search_chunks(attachment.id, "secret") == []
        assert await db.files.get_artifact("cache") is None
        assert await db.files.find_symbols(attachment.id, "f") == []
        assert await db.files.attachments_for_messages([message.id]) == {}
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_read_failed_attachment_is_not_reported_as_waiting(tmp_path):
    db = Database(str(tmp_path / "failed.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "failed.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        attachment = await db.files.create_attachment(
            session_id="session", filename="a.txt", mime_type="text/plain",
            size_bytes=1, sha256="3" * 64, storage_key="blobs/33/placeholder",
        )
        await db.files.update_attachment(attachment.id, status="failed", error_message="parse failed")

        result = await runtime.read_file("session", attachment.id)

        assert result["waiting"] is False
        assert result["error"] == "parse failed"
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_empty_ocr_image_summary_does_not_fail_attachment(tmp_path):
    db = Database(str(tmp_path / "image-summary.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "image-summary.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        attachment = await db.files.create_attachment(
            session_id="session", filename="shot.png", mime_type="image/png",
            size_bytes=1, sha256="4" * 64, storage_key="blobs/44/placeholder",
        )
        await db.files.update_attachment(attachment.id, status="ready", adapter_name="image")

        summary = await runtime.summarize_file("session", attachment.id)
        current = await db.files.get_attachment(attachment.id, "session")

        assert summary["no_text"] is True
        assert current and current.status.value == "ready"
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_image_adapter_passes_path_to_rapidocr(tmp_path, monkeypatch):
    from PIL import Image

    seen: dict[str, object] = {}

    class FakeRapidOCR:
        def __call__(self, img_content):
            seen["img_content"] = img_content
            return ([[None, "hello"]], None)

    rapidocr = types.ModuleType("rapidocr_onnxruntime")
    rapidocr.RapidOCR = FakeRapidOCR
    monkeypatch.setitem(sys.modules, "rapidocr_onnxruntime", rapidocr)

    image_path = tmp_path / "sample.png"
    Image.new("RGB", (12, 8), "white").save(image_path)

    result = await ImageAdapter().extract(_context(image_path, tmp_path), Settings(_env_file=None))

    assert seen["img_content"] == image_path
    assert result.units[0].content == "hello"
    assert result.metadata["ocr_available"] is True


@pytest.mark.asyncio
async def test_image_analysis_uses_ocr_fallback_without_vision(tmp_path, monkeypatch):
    from PIL import Image

    class FakeRapidOCR:
        def __call__(self, img_content):
            return ([[None, "invoice total 42"]], None)

    rapidocr = types.ModuleType("rapidocr_onnxruntime")
    rapidocr.RapidOCR = FakeRapidOCR
    monkeypatch.setitem(sys.modules, "rapidocr_onnxruntime", rapidocr)

    db = Database(str(tmp_path / "image-analysis.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "image-analysis.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        image_path = tmp_path / "sample.png"
        Image.new("RGB", (12, 8), "white").save(image_path)
        blob = await runtime.storage.save_stream(_chunks(image_path.read_bytes()))
        attachment = await db.files.create_attachment(
            session_id="session", filename="sample.png", mime_type="image/png",
            size_bytes=blob.size_bytes, sha256=blob.sha256, storage_key=blob.storage_key,
        )

        result = await runtime.analyze_file("session", attachment.id, "读取图片里的文字")

        assert result["can_describe_visual_content"] is False
        assert result["analysis_source"] == "ocr"
        assert result["ocr_text"] == "invoice total 42"
        assert "OCR" in result["message"]
    finally:
        await db.close()


def test_pdf_ocr_page_encodes_rendered_image_as_bytes(tmp_path, monkeypatch):
    from PIL import Image

    seen: dict[str, object] = {}

    class FakeRenderedPage:
        def to_pil(self):
            return Image.new("RGB", (2, 2), "white")

    class FakePage:
        def render(self, scale):
            seen["scale"] = scale
            return FakeRenderedPage()

    class FakePdfDocument:
        def __init__(self, path):
            seen["path"] = path

        def __getitem__(self, page_index):
            seen["page_index"] = page_index
            return FakePage()

    class FakeRapidOCR:
        def __call__(self, img_content):
            seen["img_content"] = img_content
            return ([[None, "pdf text"]], None)

    pdfium = types.ModuleType("pypdfium2")
    pdfium.PdfDocument = FakePdfDocument
    rapidocr = types.ModuleType("rapidocr_onnxruntime")
    rapidocr.RapidOCR = FakeRapidOCR
    monkeypatch.setitem(sys.modules, "pypdfium2", pdfium)
    monkeypatch.setitem(sys.modules, "rapidocr_onnxruntime", rapidocr)

    pdf_path = tmp_path / "sample.pdf"
    text = PdfAdapter._ocr_page(pdf_path, 3)

    assert text == "pdf text"
    assert seen["path"] == str(pdf_path)
    assert seen["page_index"] == 3
    assert seen["scale"] == 2
    assert isinstance(seen["img_content"], bytes)
    assert seen["img_content"].startswith(b"\x89PNG\r\n\x1a\n")


@pytest.mark.asyncio
async def test_summary_task_failure_does_not_mark_attachment_failed(tmp_path):
    db = Database(str(tmp_path / "summary-task.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "summary-task.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
            file_task_max_attempts=1,
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        worker = FileTaskWorker(runtime)
        attachment = await db.files.create_attachment(
            session_id="session", filename="empty.txt", mime_type="text/plain",
            size_bytes=1, sha256="5" * 64, storage_key="blobs/55/placeholder",
        )
        await db.files.update_attachment(attachment.id, status="ready", adapter_name="text")
        await worker.enqueue_task("session", attachment.id, FileTaskType.FILE_SUMMARY)
        task = await db.files.claim_next_task()
        assert task is not None

        await worker._execute_with_limit(task)
        current_task = await db.files.get_task(task.id)
        current_attachment = await db.files.get_attachment(attachment.id, "session")

        assert current_task and current_task.status == FileTaskStatus.FAILED
        assert current_attachment and current_attachment.status.value == "ready"
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_file_capability_governance_is_applied_to_runtime_manager(tmp_path):
    db = Database(str(tmp_path / "tools.db"))
    await db.connect()
    try:
        settings = Settings(
            _env_file=None,
            sqlite_db_path=str(tmp_path / "tools.db"),
            chromadb_path=str(tmp_path / "chroma"),
            file_storage_path=str(tmp_path / "storage"),
        )
        runtime = FileIntelligenceRuntime(db.files, _FakeLLM(), _FakeLLM(), settings=settings)
        manager = UnifiedToolManager()
        register_file_capabilities(manager, runtime)
        await _persist_registered_tool_configs(manager, db)
        await db.tools.update("read_file", enabled=False, risk_level="high", require_approval=True)

        restarted = UnifiedToolManager()
        register_file_capabilities(restarted, runtime)
        await _persist_registered_tool_configs(restarted, db)

        assert restarted.is_enabled("read_file") is False
        assert restarted.get_risk_level("read_file") == "high"
        assert restarted.require_approval("read_file") is True
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_continuation_is_claimed_only_once(tmp_path):
    db = Database(str(tmp_path / "continuation.db"))
    await db.connect()
    try:
        await db.sessions.create("session")
        attachment = await db.files.create_attachment(
            session_id="session", filename="a.txt", mime_type="text/plain",
            size_bytes=1, sha256="1" * 64, storage_key="blobs/11/placeholder",
        )
        await db.files.update_attachment(attachment.id, status="ready")
        await db.files.create_continuation(
            "session", "run", task_ids=[],
            request={"attachment_ids": [attachment.id], "message_id": "message", "user_message": "read"},
        )
        claimed = await db.files.claim_ready_continuations(attachment.id)
        assert len(claimed) == 1
        assert await db.files.claim_ready_continuations(attachment.id) == []
        await db.files.finish_continuation(claimed[0]["id"])
    finally:
        await db.close()


@pytest.mark.asyncio
async def test_archive_rejects_path_traversal(tmp_path):
    archive_path = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("../escape.txt", "bad")

    with pytest.raises(ValueError, match="路径穿越"):
        await ArchiveAdapter().extract(
            _context(archive_path, tmp_path / "workspace"),
            Settings(_env_file=None),
        )


def test_cache_key_changes_with_every_version_dimension(tmp_path):
    from datetime import datetime
    from athena.models.file import Attachment

    attachment = Attachment(
        id="f", session_id="s", filename="a.txt", mime_type="text/plain",
        size_bytes=1, sha256="a" * 64, storage_key="blobs/aa/hash",
        created_at=datetime.now(), updated_at=datetime.now(),
    )
    base = FileIntelligenceRuntime.cache_key(attachment, "summary", {"kind": "short"}, "1", "m1", "p1")
    assert base != FileIntelligenceRuntime.cache_key(attachment, "summary", {"kind": "short"}, "2", "m1", "p1")
    assert base != FileIntelligenceRuntime.cache_key(attachment, "summary", {"kind": "short"}, "1", "m2", "p1")
    assert base != FileIntelligenceRuntime.cache_key(attachment, "summary", {"kind": "short"}, "1", "m1", "p2")


@pytest.mark.asyncio
async def test_office_pdf_and_image_adapters_preserve_structure(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from pypdf import PdfWriter

    settings = Settings(_env_file=None)

    document = Document()
    document.add_heading("Title", level=1)
    document.add_paragraph("Body")
    docx_path = tmp_path / "sample.docx"
    document.save(docx_path)
    word = await WordAdapter().extract(_context(docx_path, tmp_path), settings)
    assert word.units[0].metadata["style"] == "Heading 1"

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 2
    sheet["A2"] = "=A1*2"
    xlsx_path = tmp_path / "sample.xlsx"
    workbook.save(xlsx_path)
    workbook.close()
    excel = await ExcelAdapter().extract(_context(xlsx_path, tmp_path), settings)
    assert excel.metadata["sheets"][0]["formulas"] == [{"cell": "A2", "formula": "=A1*2"}]

    image_path = tmp_path / "sample.png"
    Image.new("RGB", (12, 8), "white").save(image_path)
    image = await ImageAdapter().extract(_context(image_path, tmp_path), settings)
    assert (image.metadata["width"], image.metadata["height"]) == (12, 8)

    pdf_path = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as output:
        writer.write(output)
    pdf = await PdfAdapter().extract(_context(pdf_path, tmp_path), settings)
    assert pdf.metadata["pages"] == 1
